"""MOM JSON validation, OpenAI call settings and .docx generation (OpenAI is mocked)."""
import io
from types import SimpleNamespace

from docx import Document

from app import doc_generator, mom_generator
from app.meeting_filter import parse_vtt_transcript, transcript_to_readable, transcript_view_lines


def test_validate_keeps_only_known_shapes():
    data = {"tldr": ["a", "", 3], "action_items": [{"description": "x"}, {"assigned_to": "no desc"}, "bad"],
            "decisions": "not a list"}
    out = mom_generator._validate(data)
    assert out == {"tldr": ["a", "3"], "action_items": [
        {"description": "x", "assigned_to": "", "due_date": ""}], "decisions": []}


def test_generate_returns_error_without_api_key(monkeypatch):
    monkeypatch.setattr(mom_generator.Config, "OPENAI_API_KEY", "")
    result, error = mom_generator.generate_mom_from_transcript("hello", "S")
    assert result == mom_generator._empty() and "OPENAI_API_KEY" in error


def test_generate_parses_fenced_json(monkeypatch):
    created = {}

    class FakeClient:
        def __init__(self, **kwargs):
            created.update(kwargs)
            self.chat = SimpleNamespace(completions=SimpleNamespace(create=self._create))

        def _create(self, **kwargs):
            created["model"] = kwargs["model"]
            content = '```json\n{"tldr": ["Done"], "action_items": [], "decisions": ["Go"]}\n```'
            return SimpleNamespace(choices=[SimpleNamespace(message=SimpleNamespace(content=content))])

    monkeypatch.setattr(mom_generator.Config, "OPENAI_API_KEY", "sk-test")
    monkeypatch.setattr(mom_generator.openai, "OpenAI", FakeClient)
    mom_generator._client_cache.clear() if hasattr(mom_generator, "_client_cache") else None
    result, error = mom_generator.generate_mom_from_transcript("transcript", "S")
    assert error is None
    assert result == {"tldr": ["Done"], "action_items": [], "decisions": ["Go"]}
    assert created["api_key"] == "sk-test"


def test_docx_contains_sections():
    raw = doc_generator.generate_mom_document("Title", "2026-10-01", "10:00", ["Point"],
                                              [{"description": "Do it", "assigned_to": "Ana", "due_date": ""}],
                                              ["Decided"], transcript_text="Speaker: hi")
    doc = Document(io.BytesIO(raw))
    text = "\n".join(p.text for p in doc.paragraphs)
    assert "MINUTES OF MEETING" in text and "Point" in text and "Decided" in text
    assert "Speaker: hi" in text
    assert doc.tables[1].rows[1].cells[1].text == "Do it"


def test_vtt_parsing_round_trip():
    vtt = "WEBVTT\n\n00:00:01.000 --> 00:00:03.000\n<v Ana>Hello</v>\n\n01:02:03.500 --> 01:02:05.000\n<v Bo>Hi?</v>\n"
    entries = parse_vtt_transcript(vtt)
    assert [e["speaker"] for e in entries] == ["Ana", "Bo"]
    assert transcript_to_readable(entries) == "\nAna:\n  Hello\n\nBo:\n  Hi?"
    lines = transcript_view_lines(entries)
    assert [line["clock"] for line in lines] == ["0:01", "1:02:03"]


def test_pdf_conversion_on_linux_uses_a_per_process_profile(monkeypatch):
    seen = {}

    def fake_run(cmd, **kwargs):
        seen["cmd"] = cmd
        out_dir = cmd[cmd.index("--outdir") + 1]
        with open(f"{out_dir}/mom.pdf", "wb") as f:
            f.write(b"%PDF-1.4 fake")
        return SimpleNamespace(returncode=0, stderr="")

    monkeypatch.setattr(doc_generator.platform, "system", lambda: "Linux")
    monkeypatch.setattr(doc_generator.shutil, "which", lambda name: "/usr/bin/soffice" if name == "soffice" else None)
    monkeypatch.setattr(doc_generator.subprocess, "run", fake_run)
    assert doc_generator.convert_docx_to_pdf(b"PK fake docx") == b"%PDF-1.4 fake"
    assert seen["cmd"][0] == "/usr/bin/soffice"
    assert seen["cmd"][1].startswith("-env:UserInstallation=file:")
    assert "--convert-to" in seen["cmd"]
